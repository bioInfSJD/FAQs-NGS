#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build faqs.html from the template (faqs.template.html) and the panel
archive (Panells.zip).

You do not need to edit this file to update the panels: just replace
Panells.zip in the repository. The GitHub Action runs this script
automatically and rebuilds faqs.html.

Manual use (if you ever want to run it yourself):
    pip install openpyxl
    python generar_panells.py
"""

import os
import re
import sys
import json
import zipfile
import tempfile
import datetime
import warnings

warnings.filterwarnings("ignore")

try:
    import openpyxl
except ImportError:
    sys.exit("Missing 'openpyxl'. Install it with: pip install openpyxl")

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
ZIP_PATH = "Panells.zip"
TEMPLATE_PATH = "faqs.template.html"
OUTPUT_PATH = "faqs.html"
PLACEHOLDER = "/*__PANELLS_DATA__*/"

# Official name of each disease group (from the "Abast" sheet in
# 00.Llistat_panells.xlsx). Kept in Catalan because it is shown on the page.
ABAST = {
    'NRL': 'Neuropediatria i neurodesenvolupament',
    'GAS': 'Gastroenterologia',
    'MET': 'Malalties metabòliques i mitocondrials',
    'MUS': 'Malalties neuromusculars',
    'GEN': 'Síndromes dismòrfics',
    'ONC': 'Oncologia',
    'CAR': 'Cardiologia',
    'NEF': 'Nefrologia',
    'IMM': 'Immunologia',
    'END': 'Endocrinologia',
    'OFT': 'Oftalmologia',
    'HEM': 'Hematologia',
    'ORL': 'Otorrinolaringologia',
    'DER': 'Dermatologia',
    'OBS': 'Medicina fetal',
    'CIR': 'Cirurgia',
    'REU': 'Reumatologia',
    'PNE': 'Pneumologia',
    'OSS': 'Malalties òssies',
}

# Abbreviations for the "Origen" (source) column
ORIG = {
    'PanelApp_UK': 'UK',
    'PanelApp_AUS': 'AUS',
    'HPO': 'HPO',
    'Bibliografia': 'Bib',
    'Altres': 'Altres',
}


def abbr(origen):
    """Abbreviate a ';'-separated list of gene sources."""
    if not origen:
        return ''
    return ';'.join(ORIG.get(p.strip(), p.strip()) for p in str(origen).split(';'))


def fdate(value):
    """Format a date/datetime cell as YYYY-MM-DD, or '' if empty."""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime('%Y-%m-%d')
    return ''


def clean_name(filename):
    """Turn a file name into a display name and version.

    '01.Discapacitat_intellectual_v2.xlsx' -> ('Discapacitat intellectual', 'v2')
    """
    n = re.sub(r'\.url$', '', filename)     # SharePoint shortcut suffix
    n = re.sub(r'\.xlsx$', '', n)           # spreadsheet suffix
    n = re.sub(r'^\d+\.\s*', '', n)         # leading "01." prefix
    m = re.search(r'_(v\d+)$', n)
    ver = m.group(1) if m else ''
    n = re.sub(r'_v\d+$', '', n)
    n = n.replace('_', ' ').strip()
    return n, ver


def build_data(root):
    """Walk the extracted 'Panells' folder and return the data structure."""
    folders = sorted(
        d for d in os.listdir(root)
        if os.path.isdir(os.path.join(root, d))
        and re.match(r'\d\d\.', d)
        and not d.startswith('99')          # skip 99.OBSOLETS
    )

    groups = []
    n_panels = n_genes = 0
    for folder in folders:
        num, code = folder.split('.')[0], folder.split('.')[1]
        fpath = os.path.join(root, folder)
        panels = []
        for f in sorted(os.listdir(fpath)):
            if f.startswith('00.'):          # internal templates
                continue
            full = os.path.join(fpath, f)
            if f.endswith('.xlsx'):
                try:
                    wb = openpyxl.load_workbook(full, data_only=True, read_only=True)
                except Exception as e:
                    print(f"  WARNING: could not read {f}: {e}")
                    continue
                if 'Gens' not in wb.sheetnames:
                    wb.close()
                    continue
                rows = list(wb['Gens'].iter_rows(values_only=True))
                wb.close()
                genes = [
                    [r[0], abbr(r[1]), fdate(r[3] if len(r) > 3 else None)]
                    for r in rows[1:] if r and r[0]
                ]
                name, ver = clean_name(f)
                panels.append({'name': name, 'version': ver, 'n': len(genes), 'genes': genes})
                n_panels += 1
                n_genes += len(genes)
            elif f.endswith('.url'):         # shortcut to a panel in another group
                name, ver = clean_name(f)
                panels.append({'name': name, 'version': ver, 'alias': True})
        if not panels:
            continue
        groups.append({'num': num, 'code': code, 'name': ABAST.get(code, code), 'panels': panels})

    return {'updated': datetime.date.today().strftime('%Y-%m-%d'), 'groups': groups}, n_panels, n_genes


def find_panells_root(extract_dir):
    """Locate the folder that holds the disease groups (01.NRL, 02.GAS, ...)."""
    # Typical layout: <extract_dir>/Panells/01.NRL...
    for dirpath, dirnames, _ in os.walk(extract_dir):
        if any(re.match(r'\d\d\.[A-Z]{3}$', d) for d in dirnames):
            return dirpath
    return extract_dir


def main():
    if not os.path.exists(ZIP_PATH):
        sys.exit(f"{ZIP_PATH} not found")
    if not os.path.exists(TEMPLATE_PATH):
        sys.exit(f"{TEMPLATE_PATH} not found")

    with tempfile.TemporaryDirectory() as tmp:
        with zipfile.ZipFile(ZIP_PATH) as z:
            z.extractall(tmp)
        root = find_panells_root(tmp)
        data, n_panels, n_genes = build_data(root)

    data_js = "window.PANELLS=" + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ";"

    template = open(TEMPLATE_PATH, encoding='utf-8').read()
    if PLACEHOLDER not in template:
        sys.exit(f"Template does not contain the {PLACEHOLDER} marker")
    html = template.replace(PLACEHOLDER, data_js)

    with open(OUTPUT_PATH, 'w', encoding='utf-8') as fh:
        fh.write(html)

    print(f"OK · {len(data['groups'])} groups · {n_panels} panels · {n_genes} genes")
    print(f"Wrote {OUTPUT_PATH} ({len(html.encode('utf-8')):,} bytes)")


if __name__ == '__main__':
    main()
